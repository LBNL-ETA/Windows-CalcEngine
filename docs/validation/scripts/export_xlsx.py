"""Excel workbook with the complete Radiance-vs-engine data for every case.

usage: export_xlsx.py <runs dir> <WINDOW results dir> <wce dumps dir> <output.xlsx> [case ...]

Sheets: README (what and how), Summary (integrated values, all cases and properties, with
difference formulas), Diffuse rings and Direct rings (largest difference per incoming ring),
and one sheet per case with the 145 incoming patches: theta, phi, lambda, and for Tf, Tb,
Rf, Rb the Radiance and WCE directional-hemispherical value, direct part, diffuse part and
the difference of the diffuse parts as a formula. Legacy columns appear where WINDOW dumps
exist. Hemispherical values are SUMPRODUCT formulas over the patch columns, so the sheet
recalculates if a column is edited.
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from compare_radiance import PROPS, RING_EDGES, RING_PATCHES, klems_lambda  # noqa: E402
from summarize import LABELS, REFERENCE_FILES, engine_columns, legacy_columns, radiance_columns  # noqa: E402

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="DDE4F0")
NAMES = {"tf": "Tf", "tb": "Tb", "rf": "Rf", "rb": "Rb"}
GEOMETRY = {
    "v0": ("flat", 45, 12, "0.5 / 0.5", 0), "v1": ("flat", 45, 12, "0.8 / 0.2", 0),
    "v2": ("flat", 45, 12, "0.2 / 0.8", 0), "v3": ("curved, rise 1 mm", 45, 12, "0.5 / 0.5", 0),
    "v4": ("flat", 45, 12, "0.5 / 0.5", 0.2), "v5": ("flat", 45, 12, "0.7 / 0.2", 0.2),
    "t0": ("flat", 0, 12, "0.5 / 0.5", 0), "tm45": ("flat", -45, 12, "0.8 / 0.2", 0),
    "t80": ("flat", 80, 12, "0.5 / 0.5", 0), "wide": ("flat", 45, 16, "0.5 / 0.5", 0),
    "curved3": ("curved, rise 3 mm", 45, 12, "0.8 / 0.2", 0), "t0trans": ("flat", 0, 12, "0.7 / 0.2", 0.2),
}


def patch_angles() -> list[tuple[float, float]]:
    centres = [0.0] + [(a + b) / 2 for a, b in zip(RING_EDGES[1:], RING_EDGES[2:])]
    angles = []
    for theta, count in zip(centres, RING_PATCHES):
        angles.extend((theta, 360.0 * k / count) for k in range(count))
    return angles


def style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_font(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.font.name != FONT or (cell.font.bold and cell.fill.fgColor.rgb != "00DDE4F0"):
                cell.font = Font(name=FONT, bold=cell.font.bold)


def write_readme(wb: Workbook, cases: list[str]) -> None:
    ws = wb.active
    ws.title = "README"
    lines = [
        "Venetian blinds: Radiance reference against the WCE and legacy engines",
        "",
        "Radiance: genfmtx (WINDOW's bundled Radiance), Klems full basis, 20000 rays per patch, 12 ambient bounces;",
        "window-side incidence only, exterior incidence from the mirrored-tilt run. See docs/validation/venetian_radiance.md.",
        "WCE: Windows-CalcEngine 1.0.77, DirectionalDiffuse, 5 segments per slat, written by the unit tests (WCE_VENETIAN_DUMP_DIR).",
        "Legacy: LayerOptics through WINDOW 8.1.67 debug matrices, available for the v-cases only.",
        "",
        "Quantities (dimensionless, per incoming Klems patch, this library's patch ordering):",
        "  dir-hem  directional-hemispherical value for that incoming patch",
        "  direct   its direct-direct part: diagonal matrix element times the patch's projected solid angle lambda",
        "  diffuse  dir-hem minus direct",
        "  hemispherical (Summary sheet): SUMPRODUCT(dir-hem, lambda) / PI, i.e. the diffuse-diffuse value",
        "",
        "Slats 16 mm wide in every case. Cases:",
    ]
    for i, text in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text)
    ws.cell(row=1, column=1).font = Font(name=FONT, bold=True, size=12)
    start = len(lines) + 1
    headers = ["Case", "Slat shape", "Tilt (deg)", "Spacing (mm)", "Rf / Rb", "T", "Legacy data"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start, column=j, value=h)
    style_header(ws, start, len(headers))
    for i, case in enumerate(cases, start=start + 1):
        shape, tilt, spacing, refl, tau = GEOMETRY[case]
        for j, v in enumerate([case, shape, tilt, spacing, refl, tau, "yes" if case.startswith("v") else "no"], start=1):
            ws.cell(row=i, column=j, value=v)
    ws.column_dimensions["A"].width = 14
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 16
    set_font(ws)


def write_case_sheet(wb: Workbook, case: str, rad, wce, leg, lam: np.ndarray) -> dict[str, dict[str, str]]:
    """Write one case; return cell ranges of the dir-hem columns for the Summary formulas."""
    ws = wb.create_sheet(case)
    angles = patch_angles()
    header = ["Patch", "Theta (deg)", "Phi (deg)", "Lambda"]
    sources = [("Radiance", rad), ("WCE", wce)] + ([("Legacy", leg)] if leg is not None else [])
    for p in PROPS:
        for name, _ in sources:
            header += [f"{NAMES[p]} {name} dir-hem", f"{NAMES[p]} {name} direct", f"{NAMES[p]} {name} diffuse"]
        header += [f"{NAMES[p]} WCE-Rad diffuse"]
    ws.append(header)
    style_header(ws, 1, len(header))
    ranges: dict[str, dict[str, str]] = {}
    col = 5
    layout: dict[str, dict[str, int]] = {}
    for p in PROPS:
        layout[p] = {}
        for name, _ in sources:
            layout[p][name] = col
            col += 3
        layout[p]["diff"] = col
        col += 1
    for i in range(145):
        row = i + 2
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=angles[i][0])
        ws.cell(row=row, column=3, value=round(angles[i][1], 2))
        ws.cell(row=row, column=4, value=float(lam[i]))
        for p in PROPS:
            for name, cols in sources:
                c = layout[p][name]
                ws.cell(row=row, column=c, value=float(cols.dirhem[p][i]))
                ws.cell(row=row, column=c + 1, value=float(cols.direct[p][i]))
                ws.cell(row=row, column=c + 2, value=float(cols.dirhem[p][i] - cols.direct[p][i]))
            r_diff = get_column_letter(layout[p]["Radiance"] + 2)
            w_diff = get_column_letter(layout[p]["WCE"] + 2)
            ws.cell(row=row, column=layout[p]["diff"], value=f"={w_diff}{row}-{r_diff}{row}")
    for p in PROPS:
        ranges[p] = {name: f"'{case}'!${get_column_letter(layout[p][name])}$2:${get_column_letter(layout[p][name])}$146"
                     for name, _ in sources}
    for c in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = "E2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "0.000000" if cell.column >= 4 else "0.##"
    set_font(ws)
    return ranges


def write_summary(wb: Workbook, cases: list[str], ranges: dict[str, dict], data: dict) -> None:
    ws = wb.create_sheet("Summary", 1)
    ws["A1"] = "Integrated values per case and property. Normal incidence = dir-hem of patch 0; hemispherical = SUMPRODUCT(dir-hem, lambda)/PI over the case sheet. Differences are formulas."
    ws["A1"].font = Font(name=FONT, italic=True)
    header = ["Case", "Property", "Rad. normal", "WCE normal", "WCE - Rad.", "Rad. hemisph.", "WCE hemisph.", "WCE - Rad.",
              "Legacy normal", "Legacy - Rad.", "Legacy hemisph.", "Legacy - Rad."]
    ws.append([])
    ws.append(header)
    style_header(ws, 3, len(header))
    row = 4
    for case in cases:
        rad, wce, leg = data[case]
        for p in PROPS:
            rng = ranges[case][p]
            lam_range = f"'{case}'!$D$2:$D$146"
            ws.cell(row=row, column=1, value=case)
            ws.cell(row=row, column=2, value=NAMES[p])
            ws.cell(row=row, column=3, value=f"=INDEX({rng['Radiance']},1)")
            ws.cell(row=row, column=4, value=f"=INDEX({rng['WCE']},1)")
            ws.cell(row=row, column=5, value=f"=D{row}-C{row}")
            ws.cell(row=row, column=6, value=f"=SUMPRODUCT({rng['Radiance']},{lam_range})/PI()")
            ws.cell(row=row, column=7, value=f"=SUMPRODUCT({rng['WCE']},{lam_range})/PI()")
            ws.cell(row=row, column=8, value=f"=G{row}-F{row}")
            if "Legacy" in rng:
                ws.cell(row=row, column=9, value=f"=INDEX({rng['Legacy']},1)")
                ws.cell(row=row, column=10, value=f"=I{row}-C{row}")
                ws.cell(row=row, column=11, value=f"=SUMPRODUCT({rng['Legacy']},{lam_range})/PI()")
                ws.cell(row=row, column=12, value=f"=K{row}-F{row}")
            row += 1
    for r in ws.iter_rows(min_row=4, max_row=row - 1, min_col=3, max_col=12):
        for cell in r:
            cell.number_format = "0.0000;-0.0000;0.0000"
    for c, w in zip("ABCDEFGHIJKL", (10, 9, 12, 12, 12, 13, 13, 12, 13, 13, 14, 13)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "C4"
    set_font(ws)


def write_ring_sheet(wb: Workbook, title: str, cases: list[str], data: dict, part: str) -> None:
    ws = wb.create_sheet(title)
    centres = [0.0] + [(a + b) / 2 for a, b in zip(RING_EDGES[1:], RING_EDGES[2:])]
    ring = np.repeat(np.arange(len(RING_PATCHES)), RING_PATCHES)
    ws["A1"] = f"Largest |WCE - Radiance| of the {part} part per incoming ring (ring centre theta in degrees). Values, computed by export_xlsx.py from the case sheets."
    ws["A1"].font = Font(name=FONT, italic=True)
    ws.append([])
    ws.append(["Case", "Property"] + [f"{c:g}" for c in centres])
    style_header(ws, 3, 2 + len(centres))
    for case in cases:
        rad, wce, _ = data[case]
        for p in PROPS:
            if part == "diffuse":
                diff = np.abs(rad.diffuse(p) - wce.diffuse(p))
            else:
                diff = np.abs(rad.direct[p] - wce.direct[p])
            ws.append([case, NAMES[p]] + [float(diff[ring == r].max()) for r in range(len(RING_PATCHES))])
    for r in ws.iter_rows(min_row=4, min_col=3):
        for cell in r:
            cell.number_format = "0.000"
    ws.column_dimensions["A"].width = 10
    ws.freeze_panes = "C4"
    set_font(ws)


def main() -> int:
    runs, results, dumps, out = (Path(a) for a in sys.argv[1:5])
    cases = sys.argv[5:] or list(REFERENCE_FILES)
    lam = klems_lambda()
    data = {}
    for case in cases:
        wce = engine_columns(results, dumps, case, lam)
        if wce is None:
            print(f"skip {case}: no WCE data")
            continue
        data[case] = (radiance_columns(runs, case, lam), wce, legacy_columns(results, case, lam))
    cases = list(data)

    wb = Workbook()
    write_readme(wb, cases)
    ranges = {}
    for case in cases:
        rad, wce, leg = data[case]
        ranges[case] = write_case_sheet(wb, case, rad, wce, leg, lam)
    write_summary(wb, cases, ranges, data)
    write_ring_sheet(wb, "Diffuse rings", cases, data, "diffuse")
    write_ring_sheet(wb, "Direct rings", cases, data, "direct")
    wb.move_sheet("Diffuse rings", offset=-(len(cases)))
    wb.move_sheet("Direct rings", offset=-(len(cases)))
    wb.calculation.fullCalcOnLoad = True   # openpyxl stores no cached values; Excel computes on open
    wb.save(out)
    print(f"{out}: {len(cases)} cases, sheets {wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
